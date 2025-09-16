import os
import shutil
from datetime import datetime

import docx
from airflow import DAG
from airflow.providers.postgres.hooks.postgres import PostgresHook
from airflow.providers.standard.operators.python import PythonOperator
from dotenv import load_dotenv

WATCH_DIR = "/opt/airflow/data/incoming"


def list_subfolder(**context):
    """Find subfolder passed by trigger"""
    conf = context["dag_run"].conf
    subfolder = conf.get("subfolder")  # folder name represents the name of the group
    if not subfolder:
        raise ValueError('You must trigger DAG with {"subfolder": "GroupName"}')
    folder_path = os.path.join(WATCH_DIR, subfolder)
    if not os.path.isdir(folder_path):
        raise FileNotFoundError(f"Subfolder {folder_path} not found")

    files = [
        os.path.join(folder_path, f)
        for f in os.listdir(folder_path)
        if f.endswith(".docx")
    ]
    context["ti"].xcom_push(key="group_name", value=subfolder)
    context["ti"].xcom_push(key="files", value=files)


def parse_docs(**context):
    files = context["ti"].xcom_pull(key="files", task_ids="list_subfolder")
    group_name = context["ti"].xcom_pull(key="group_name", task_ids="list_subfolder")

    students = []
    for file in files:
        fname = os.path.basename(file).replace(".docx", "")
        try:
            first, last = fname.split("_", 1)
        except ValueError:
            raise ValueError(
                f"Invalid filename format: {fname}. Expected Name_Surname.docx"
            )

        doc = docx.Document(file)
        text = "\n".join([p.text for p in doc.paragraphs])
        students.append(
            {"group": group_name, "first": first, "last": last, "text": text}
        )

    context["ti"].xcom_push(key="students", value=students)


def insert_postgres(**context):
    conf = context["dag_run"].conf
    students = context["ti"].xcom_pull(key="students", task_ids="parse_docs")
    graduation_date = conf.get("graduation_date")
    courses = conf.get("courses")
    group_name = conf.get("subfolder")  # folder name represents the name of the group

    if not graduation_date:
        raise ValueError('You must trigger DAG with {"graduation_date": "YYYY-MM-DD"}')
    if not courses:
        raise ValueError(
            'You must trigger DAG with {"courses": ["Subject_name_1", "Subject_name_2", ...]}'
        )

    load_dotenv()
    pg_hook = PostgresHook(
        host="postgres",
        schema=os.getenv("POSTGRES_MAIN_DB"),
        login=os.getenv("POSTGRES_USER"),
        password=os.getenv("POSTGRES_PASSWORD"),
        port=5432,
    )

    insert_groups_sql = """
        INSERT INTO groups (code, graduation_date)
        VALUES (%s, %s)
    """
    insert_students_sql = """
        INSERT INTO students (group_code, first_name, last_name, essay_text)
        VALUES (%s, %s, %s, %s)
    """
    placeholders = ", ".join(
        ["%s"] * len(courses)
    )  # creating placeholders list like [%s, %s, ...] to prevent sql injection
    select_courses_ids_sql = f"SELECT id FROM courses WHERE name IN ({placeholders})"
    insert_groups_courses_sql = """
        INSERT INTO groups_courses(group_code, course_id) VALUES (%s, %s)
        ON CONFLICT DO NOTHING;
    """

    pg_hook.run(insert_groups_sql, parameters=(students[0]["group"], graduation_date))
    for student in students:
        pg_hook.run(
            insert_students_sql,
            parameters=(
                student["group"],
                student["first"],
                student["last"],
                student["text"],
            ),
        )

    courses_ids = pg_hook.get_records(select_courses_ids_sql, parameters=courses)
    for (course_id,) in courses_ids:
        pg_hook.run(
            insert_groups_courses_sql,
            parameters=(group_name, course_id),
        )


def cleanup_folder(**context):
    group_name = context["ti"].xcom_pull(key="group_name", task_ids="list_subfolder")
    folder_path = os.path.join(WATCH_DIR, group_name)
    if os.path.isdir(folder_path):
        shutil.rmtree(folder_path)


with DAG(
    dag_id="process_student_docs",
    start_date=datetime(2025, 9, 14),
    schedule=None,
    catchup=False,
) as dag:
    t1 = PythonOperator(
        task_id="list_subfolder",
        python_callable=list_subfolder,
    )

    t2 = PythonOperator(
        task_id="parse_docs",
        python_callable=parse_docs,
    )

    t3 = PythonOperator(
        task_id="insert_postgres",
        python_callable=insert_postgres,
    )

    t4 = PythonOperator(
        task_id="cleanup_folder",
        python_callable=cleanup_folder,
    )

    t1 >> t2 >> t3 >> t4  # type:ignore
